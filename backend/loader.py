#!/usr/bin/env python3

import json
import urllib.request
from pathlib import Path

from ocr import is_image_path, ocr_image_bytes, ocr_pdf_page


class Inbox:
    def __init__(self, source):
        self.source = source.rstrip("/")
        self.is_http = (
            self.source.startswith("http://")
            or self.source.startswith("https://")
        )

    # -- listing ---------------------------------------------------------

    def emails(self):
        """Return the list of email records (dicts)."""
        if self.is_http:
            return self._get_json("/emails")

        inbox_dir = Path(self.source) / "inbox"

        return [
            json.loads(p.read_text())
            for p in sorted(inbox_dir.glob("email_*.json"))
        ]

    def __iter__(self):
        return iter(self.emails())

    def get(self, email_id):
        if self.is_http:
            return self._get_json(f"/emails/{email_id}")

        return json.loads(
            (Path(self.source) / "inbox" / f"{email_id}.json").read_text()
        )

    # -- attachments -----------------------------------------------------

    def read_bytes(self, att_path):
        """Read raw bytes from an attachment."""

        if self.is_http:
            return self._get_bytes("/" + att_path.lstrip("/"))

        return (Path(self.source) / att_path).read_bytes()

    def read_text(self, att_path, encoding="utf-8"):
        """Read an attachment as text, including PDF and Excel files."""

        return self.read_document(att_path, encoding)["text"]

    def read_document(self, att_path, encoding="utf-8"):
        """Read an attachment and report how the text was obtained.

        Returns {"text": str, "ocr": bool}. "ocr" is True when any part
        of the text came from OCR (a scanned page or an image file), so
        callers can treat it as less trustworthy than a real text layer.
        """

        data = self.read_bytes(att_path)

        # Image attachments (scans, photos)
        if is_image_path(att_path):
            return {
                "text": ocr_image_bytes(data),
                "ocr": True,
            }

        # PDF
        if att_path.lower().endswith(".pdf"):
            from io import BytesIO
            from pypdf import PdfReader

            reader = PdfReader(BytesIO(data))

            pages = []
            used_ocr = False

            for page in reader.pages:
                text = page.extract_text()

                if not text or not text.strip():
                    # Scanned page: no text layer, so OCR its images
                    text = ocr_pdf_page(page)

                    if text:
                        used_ocr = True

                if text:
                    pages.append(text)

            return {
                "text": "\n".join(pages),
                "ocr": used_ocr,
            }

        # Excel
        if att_path.lower().endswith((".xlsx", ".xlsm")):
            from io import BytesIO
            from openpyxl import load_workbook

            workbook = load_workbook(
                filename=BytesIO(data),
                data_only=True
            )

            lines = []

            for sheet in workbook.worksheets:
                lines.append(f"--- SHEET: {sheet.title} ---")

                for row in sheet.iter_rows(values_only=True):
                    values = []

                    for value in row:
                        if value is not None:
                            values.append(str(value))

                    if values:
                        lines.append(" | ".join(values))

            return {
                "text": "\n".join(lines),
                "ocr": False,
            }

        # Normal text files
        return {
            "text": data.decode(encoding, errors="replace"),
            "ocr": False,
        }

    # -- submission ------------------------------------------------------

    def submit(self, submission):
        """POST a submission to the server and return the scoreboard."""

        if not self.is_http:
            raise RuntimeError(
                "submit() needs an HTTP source; run the docker server"
            )

        data = json.dumps(submission).encode()

        req = urllib.request.Request(
            self.source + "/submit",
            data=data,
            headers={"Content-Type": "application/json"},
        )

        with urllib.request.urlopen(req) as r:
            return json.loads(r.read())

    def sample_submission(self):
        if self.is_http:
            return self._get_json("/sample_submission")

        return json.loads(
            (Path(self.source) / "sample_submission.json").read_text()
        )

    # -- HTTP helpers ----------------------------------------------------

    def _get_json(self, path):
        with urllib.request.urlopen(self.source + path) as r:
            return json.loads(r.read())

    def _get_bytes(self, path):
        with urllib.request.urlopen(self.source + path) as r:
            return r.read()


if __name__ == "__main__":
    import sys

    src = sys.argv[1] if len(sys.argv) > 1 else "data"

    inbox = Inbox(src)

    ems = inbox.emails()

    print(f"{len(ems)} emails from {src}")

    docs = [e for e in ems if e["attachments"]]

    print(
        f"{len(docs)} have attachments; "
        f"example: {docs[0]['email_id']}"
    )

    for a in docs[0]["attachments"]:
        if a.lower().endswith(".txt"):
            head = (
                inbox.read_text(a)[:60]
                .replace("\n", " ")
            )
        else:
            head = "(binary)"

        print(f"  {a}: {head}")